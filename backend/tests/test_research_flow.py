"""End-to-end research workflow through the API, including every specification quality gate."""

from __future__ import annotations

import io

from conftest import SAMPLES
from openpyxl import load_workbook


def adopt_generated_design(project_id: str) -> None:
    """Adopt the deterministic design package directly (the agent path is tested separately)."""
    from mdos.db import session_factory
    from mdos.models import Project
    from mdos.research.design import generate_all
    from mdos.services.research import adopt_design

    db = session_factory()()
    try:
        project = db.get(Project, project_id)
        package = generate_all({"name": project.name, "business_question": project.business_question,
                                "industry": project.industry, "currency": project.currency})
        adopt_design(db, project, package, run_id="test-run", user_id=project.created_by)
        db.commit()
    finally:
        db.close()


def upload_survey(client, pid: str) -> dict:
    with open(SAMPLES / "lake_toba_survey.csv", "rb") as f:
        res = client.post(f"/api/v1/projects/{pid}/datasets", files={"file": ("lake_toba_survey.csv", f, "text/csv")},
                          data={"name": "Lake Toba survey", "kind": "survey"})
    assert res.status_code == 201, res.text
    return res.json()


def test_full_research_flow(local_client, project):
    pid = project["id"]
    c = local_client
    adopt_generated_design(pid)
    overview = c.get(f"/api/v1/projects/{pid}/research").json()
    assert len(overview["hypotheses"]) == 5 and len(overview["constructs"]) == 5
    assert {"framing", "design", "sampling", "fieldwork"} <= set(overview["plans"])
    survey_id = overview["surveys"][0]["id"]

    # Questionnaire exports
    xls = c.get(f"/api/v1/projects/{pid}/surveys/{survey_id}/export?format=xlsform")
    assert xls.status_code == 200
    wb = load_workbook(io.BytesIO(xls.content))
    assert {"survey", "choices", "settings"} <= set(wb.sheetnames)
    names = [row[1] for row in wb["survey"].iter_rows(values_only=True)]
    assert "wtp_150k" in names and "vw_too_expensive" in names
    md = c.get(f"/api/v1/projects/{pid}/surveys/{survey_id}/export?format=markdown").text
    assert "`wtp_150k`" in md and "Sangat setuju" in md  # bilingual printable questionnaire
    codebook = c.get(f"/api/v1/projects/{pid}/surveys/{survey_id}/export?format=codebook").text
    assert "wtp_150k" in codebook and "5=Strongly agree" in codebook

    # Upload and quality diagnostics (variables from the design drive attention checks and scales)
    dataset = upload_survey(c, pid)
    v1 = dataset["versions"][0]
    quality = c.get(f"/api/v1/projects/{pid}/datasets/{dataset['id']}/versions/{v1['id']}").json()["quality"]
    checks = {i["check"] for i in quality["issues"]}
    assert {"attention_check", "straightlining", "speeders", "duplicate_rows", "personal_data", "price_inconsistency"} <= checks

    # Cleaning plan requires approval, then creates version 2 with lineage
    plan = c.post(f"/api/v1/projects/{pid}/datasets/{dataset['id']}/cleaning-plan").json()
    assert plan["operations"] and plan["approval_id"]
    pending = c.get(f"/api/v1/projects/{pid}/approvals").json()
    assert any(a["action"] == "apply_cleaning" for a in pending)
    decided = c.post(f"/api/v1/projects/{pid}/approvals/{plan['approval_id']}/decide", json={"decision": "approved"}).json()
    assert decided["outcome"]["applied"] and decided["outcome"]["rows"] < 326
    scaled = c.post(f"/api/v1/projects/{pid}/datasets/{dataset['id']}/compute-scales")
    assert scaled.status_code == 201, scaled.text
    v3 = scaled.json()
    assert {"CI", "AU", "PV", "PC", "PI"} <= {col["name"] for col in v3["columns"]}
    assert v3["parent_id"] and v3["operations"][0]["op"] == "compute_scale"

    # Analyses produce assumption checks and evidence candidates
    vw = c.post(f"/api/v1/projects/{pid}/analyses", json={
        "dataset_version_id": v3["id"], "method": "van_westendorp",
        "params": {"too_cheap": "vw_too_cheap", "cheap": "vw_cheap", "expensive": "vw_expensive",
                   "too_expensive": "vw_too_expensive", "target_price": 150000}}).json()
    assert vw["assumptions"] and vw["result"]["data"]["pmc"] < vw["result"]["data"]["pme"]
    ols = c.post(f"/api/v1/projects/{pid}/analyses", json={
        "dataset_version_id": v3["id"], "method": "regression_ols", "params": {"dv": "PI", "predictors": ["PV", "CI", "PC"]}}).json()
    assert len(ols["assumptions"]) >= 6
    ev = c.post(f"/api/v1/projects/{pid}/analyses/{ols['id']}/evidence", json={"keys": ["ols:PI:PV"]}).json()
    assert ev[0]["code"] == "E1" and ev[0]["origin"] == "user_data"
    ev_vw = c.post(f"/api/v1/projects/{pid}/analyses/{vw['id']}/evidence", json={}).json()
    assert [e["code"] for e in ev_vw] == ["E2", "E3"]
    # Evidence values are copied from the analysis, not re-derived
    candidate = next(x for x in ols["result"]["evidence_candidates"] if x["key"] == "ols:PI:PV")
    assert ev[0]["p_value"] == candidate["p_value"] and ev[0]["statement"] == candidate["statement"]

    # Quality gate: segmentation needs an objective and rationale
    seg = c.post(f"/api/v1/projects/{pid}/analyses", json={
        "dataset_version_id": v3["id"], "method": "segmentation", "params": {"variables": ["PV", "PC"]}})
    assert seg.status_code == 422

    # Quality gate: causal language on survey evidence is rejected, associational wording passes
    bad = c.post(f"/api/v1/projects/{pid}/insights", json={
        "title": "Value drives intention", "statement": "Higher perceived value causes more bookings.", "evidence_ids": [ev[0]["id"]]})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "quality_gate"
    assert "suggestion" in bad.json()["error"]["details"]
    good = c.post(f"/api/v1/projects/{pid}/insights", json={
        "title": "Value and intention move together", "statement": "Tourists who see more value report higher intention to book.",
        "evidence_ids": [ev[0]["id"]]})
    assert good.status_code == 201
    insight = good.json()
    assert insight["status"] == "draft"
    # Quality gate: an insight without evidence is rejected
    assert c.post(f"/api/v1/projects/{pid}/insights", json={"title": "No evidence", "statement": "Something is true here."}).status_code == 422
    approved = c.post(f"/api/v1/projects/{pid}/insights/{insight['id']}/decide", json={"decision": "approved"}).json()
    assert approved["status"] == "approved"

    # Quality gate: every recommendation cites evidence
    assert c.post(f"/api/v1/projects/{pid}/recommendations", json={"statement": "Launch at Rp 150,000."}).status_code == 422
    rec = c.post(f"/api/v1/projects/{pid}/recommendations", json={
        "statement": "Launch at Rp 150,000 and validate with a pre-sale test.", "evidence_ids": [ev_vw[1]["id"]]})
    assert rec.status_code == 201

    # Quality gate: verdicts are proposals until a person approves them with evidence
    h1 = next(h for h in overview["hypotheses"] if h["code"] == "H1")
    proposed = c.post(f"/api/v1/projects/{pid}/hypotheses/{h1['id']}/verdict",
                      json={"verdict": "supported", "evidence_ids": [ev[0]["id"]], "rationale": "b > 0, p < .001"}).json()
    assert proposed["status"] == "proposed_supported"
    final = c.post(f"/api/v1/projects/{pid}/hypotheses/{h1['id']}/verdict/decide", json={"approve": True}).json()
    assert final["status"] == "supported"

    # Integrity: evidence in use and analyses with evidence cannot be deleted
    assert c.delete(f"/api/v1/projects/{pid}/evidence/{ev[0]['id']}").status_code == 409
    assert c.delete(f"/api/v1/projects/{pid}/analyses/{ols['id']}").status_code == 409

    # Evidence graph links data to decisions
    graph = c.get(f"/api/v1/projects/{pid}/evidence-graph").json()
    kinds = {n["kind"] for n in graph["nodes"]}
    assert {"dataset", "dataset_version", "analysis", "evidence", "insight", "recommendation", "hypothesis"} <= kinds
    assert graph["edges"]

    # Report cites evidence codes and exports to Markdown and HTML
    report = c.post(f"/api/v1/projects/{pid}/reports", json={"kind": "research_report"}).json()
    md = c.get(f"/api/v1/projects/{pid}/reports/{report['id']}/export?format=md").text
    assert "[E1]" in md and "Evidence register" in md and "Methods appendix" in md
    html = c.get(f"/api/v1/projects/{pid}/reports/{report['id']}/export?format=html").text
    assert 'id="E1"' in html and "<script" not in html
    final_report = c.post(f"/api/v1/projects/{pid}/reports/{report['id']}/finalize").json()
    assert final_report["status"] == "final"

    progress = {s["name"]: s["status"] for s in c.get(f"/api/v1/projects/{pid}/research/progress").json()}
    assert progress["Data cleaning"] == "done" and progress["Report generation"] == "done"
    assert progress["Presentation generation"] == "deferred"


def test_upload_limits_and_formula_injection(make_client):
    c = make_client("local", MAX_UPLOAD_MB="1")
    token = c.post("/api/v1/auth/local-session").json()["access_token"]
    c.headers.update({"Authorization": f"Bearer {token}"})
    pid = c.post("/api/v1/projects", json={"name": "P", "business_question": "Would customers pay for this?"}).json()["id"]
    big = b"a,b\n" + b"1,2\n" * 400_000
    assert c.post(f"/api/v1/projects/{pid}/datasets", files={"file": ("big.csv", big, "text/csv")}).status_code == 422
    assert c.post(f"/api/v1/projects/{pid}/datasets", files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 422
    semicolon = b"id;rating;comment\n1;5;=HYPERLINK(\"http://evil\")\n2;4;ok\n3;3;fine\n"
    res = c.post(f"/api/v1/projects/{pid}/datasets", files={"file": ("semi.csv", semicolon, "text/csv")})
    assert res.status_code == 201
    assert {col["name"] for col in res.json()["versions"][0]["columns"]} == {"id", "rating", "comment"}

    from mdos.research.questionnaire import safe_cell

    assert safe_cell("=1+1") == "'=1+1" and safe_cell("hello") == "hello"


def test_viewer_cannot_mutate(make_client):
    c = make_client("cloud")
    owner = c.post("/api/v1/auth/register", json={"email": "own@example.com", "password": "correct horse battery",
                                                  "name": "Owner", "organization": "Org"}).json()
    headers = {"Authorization": f"Bearer {owner['access_token']}"}
    pid = c.post("/api/v1/projects", json={"name": "P", "business_question": "Would tourists pay for a tour?"}, headers=headers).json()["id"]
    # A second member of the same organization, added to the project as viewer.
    from mdos.db import session_factory
    from mdos.models import User
    from mdos.security import create_access_token, hash_password

    db = session_factory()()
    viewer = User(org_id=owner["user"]["org_id"], email="view@example.com", name="Viewer",
                  password_hash=hash_password("correct horse battery"), role="member")
    db.add(viewer)
    db.commit()
    viewer_token = create_access_token(viewer.id, viewer.org_id)
    db.close()
    added = c.post(f"/api/v1/projects/{pid}/members", json={"email": "view@example.com", "role": "viewer"}, headers=headers)
    assert added.status_code == 201
    vh = {"Authorization": f"Bearer {viewer_token}"}
    assert c.get(f"/api/v1/projects/{pid}", headers=vh).status_code == 200
    assert c.patch(f"/api/v1/projects/{pid}", json={"name": "Hacked"}, headers=vh).status_code == 403
    assert c.post(f"/api/v1/projects/{pid}/hypotheses", json={"statement": "X is associated with Y"}, headers=vh).status_code == 403
