"""Questionnaire exports: XLSForm (KoboToolbox, ODK, SurveyCTO), Markdown and a codebook CSV."""

from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import Workbook

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    """Neutralize spreadsheet formula injection in exported text cells."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def _list_name(code: str) -> str:
    return re.sub(r"[^a-z0-9_]", "_", code.lower()) + "_opts"


def to_xlsform(survey: dict[str, Any], questions: list[dict[str, Any]]) -> bytes:
    """Build an XLSForm workbook. Likert items share one 'agree5' choice list."""
    wb = Workbook()
    ws = wb.active
    ws.title = "survey"
    ws.append(["type", "name", "label::English (en)", "label::Indonesian (id)", "required", "relevant", "constraint",
               "constraint_message::English (en)"])
    choices = wb.create_sheet("choices")
    choices.append(["list_name", "name", "label::English (en)", "label::Indonesian (id)"])
    settings = wb.create_sheet("settings")
    settings.append(["form_title", "form_id", "default_language", "version"])
    form_id = re.sub(r"[^a-z0-9_]", "_", (survey.get("title") or "mdos_survey").lower())[:40] or "mdos_survey"
    settings.append([safe_cell(survey.get("title", "MDOS survey")), form_id, "English (en)", str(survey.get("version", 1))])

    ws.append(["note", "consent_note", safe_cell(survey.get("consent_text", "")), "", "", "", "", ""])
    agree_added = False
    screen_relevant = ""
    current_section = None
    for q in sorted(questions, key=lambda x: x.get("position", 0)):
        section = q.get("section") or ""
        if section != current_section:
            if current_section is not None:
                ws.append(["end_group", "", "", "", "", "", "", ""])
            group_name = re.sub(r"[^a-z0-9_]", "_", section.lower()) or "section"
            ws.append(["begin_group", f"grp_{group_name}", safe_cell(section), "", "",
                       screen_relevant if section != "Screening" else "", "", ""])
            current_section = section
        qtype, constraint, message = q["qtype"], "", ""
        required = "yes" if q.get("required", True) else ""
        if qtype == "info":
            xtype, required = "note", ""
        elif qtype == "likert":
            scale = q.get("scale") or {}
            if int(scale.get("max", 5)) == 5 and int(scale.get("min", 1)) == 1:
                xtype = "select_one agree5"
                if not agree_added:
                    labels_en = scale.get("labels") or ["1", "2", "3", "4", "5"]
                    labels_id = scale.get("labels_id") or labels_en
                    for i in range(5):
                        choices.append(["agree5", str(i + 1), labels_en[i], labels_id[i]])
                    agree_added = True
            else:
                xtype = "integer"
                constraint = f". >= {scale.get('min', 1)} and . <= {scale.get('max', 7)}"
                message = f"Enter a value from {scale.get('min', 1)} to {scale.get('max', 7)}"
        elif qtype in ("single", "multi"):
            list_name = _list_name(q["code"])
            xtype = f"{'select_one' if qtype == 'single' else 'select_multiple'} {list_name}"
            for opt in q.get("options", []):
                choices.append([list_name, safe_cell(str(opt["value"])), safe_cell(opt.get("label", opt["value"])),
                                safe_cell(opt.get("label_id", opt.get("label", opt["value"])))])
        elif qtype == "numeric":
            xtype = "integer"
            scale = q.get("scale") or {}
            if "min" in scale and "max" in scale:
                constraint = f". >= {scale['min']} and . <= {scale['max']}"
                message = f"Enter a value from {scale['min']} to {scale['max']}"
        elif qtype == "price":
            xtype, constraint, message = "integer", ". >= 0", "Enter a price of zero or more"
        else:
            xtype = "text"
        ws.append([xtype, q["code"], safe_cell(q["text"]), safe_cell(q.get("text_id", "")), required, "", constraint, message])
        terminate = (q.get("logic") or {}).get("terminate_if")
        if terminate is not None:
            screen_relevant = f"${{{q['code']}}} != '{terminate}'"
    if current_section is not None:
        ws.append(["end_group", "", "", "", "", "", "", ""])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def to_markdown(survey: dict[str, Any], questions: list[dict[str, Any]], bilingual: bool = True) -> str:
    lines = [f"# {survey.get('title', 'Questionnaire')}", ""]
    if survey.get("introduction"):
        lines += [survey["introduction"], ""]
    if survey.get("consent_text"):
        lines += [f"> {survey['consent_text']}", ""]
    section = None
    number = 0
    for q in sorted(questions, key=lambda x: x.get("position", 0)):
        if q.get("section") != section:
            section = q.get("section")
            lines += [f"## {section}", ""]
        if q["qtype"] == "info":
            lines += [f"_{q['text']}_", ""]
            continue
        number += 1
        lines.append(f"**Q{number}. {q['text']}** `{q['code']}`")
        if bilingual and q.get("text_id"):
            lines.append(f"_{q['text_id']}_")
        if q["qtype"] == "likert":
            scale = q.get("scale") or {}
            labels = scale.get("labels") or []
            lines.append("  " + " · ".join(f"{i + int(scale.get('min', 1))} {lab}" for i, lab in enumerate(labels)))
            if bilingual and scale.get("labels_id"):
                lines.append("  _" + " · ".join(f"{i + int(scale.get('min', 1))} {lab}"
                                                for i, lab in enumerate(scale["labels_id"])) + "_")
        elif q.get("options"):
            for opt in q["options"]:
                label = opt.get("label", opt["value"])
                label_id = opt.get("label_id")
                lines.append(f"  - [ ] {label}" + (f" / {label_id}" if bilingual and label_id and label_id != label else ""))
        elif q["qtype"] == "price":
            lines.append(f"  Amount ({(q.get('scale') or {}).get('currency', '')}): ________")
        elif q["qtype"] == "numeric":
            lines.append("  Number: ________")
        else:
            lines.append("  ______________________________________________")
        logic = q.get("logic") or {}
        if logic.get("terminate_if"):
            lines.append(f"  _If '{logic['terminate_if']}', thank the respondent and end the survey._")
        lines.append("")
    return "\n".join(lines)


def to_codebook_csv(questions: list[dict[str, Any]], variables: dict[str, dict[str, Any]] | None = None) -> str:
    variables = variables or {}
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["variable", "question", "question_id", "type", "values", "construct", "role", "section"])
    for q in sorted(questions, key=lambda x: x.get("position", 0)):
        if q["qtype"] == "info":
            continue
        if q["qtype"] == "likert":
            scale = q.get("scale") or {}
            values = "; ".join(f"{i + int(scale.get('min', 1))}={lab}" for i, lab in enumerate(scale.get("labels") or []))
        elif q.get("options"):
            values = "; ".join(str(o["value"]) for o in q["options"])
        else:
            values = ""
        var = variables.get(q["code"], {})
        writer.writerow([safe_cell(v) for v in (q["code"], q["text"], q.get("text_id", ""), q["qtype"], values,
                                                  q.get("construct_code", ""), var.get("role", ""), q.get("section", ""))])
    return buffer.getvalue()
